import codecs
import contextlib
import zipfile
from pathlib import Path
from unittest.mock import MagicMock

import docx
import pytest

from quickfilelocator.readers import (
    DETECTION_SAMPLE_BYTES,
    DocxFileReader,
    TextFileReader,
    UnreadableFileError,
    default_readers,
)

JAPANESE_TEXT = "日本語の文書を検索するテストです。"


@pytest.fixture
def detect():
    return MagicMock(return_value={"encoding": None})


@pytest.fixture
def text_reader(detect):
    return TextFileReader(detect=detect)


@pytest.mark.parametrize(
    "data",
    [
        codecs.BOM_UTF8 + JAPANESE_TEXT.encode("utf-8"),
        codecs.BOM_UTF16_LE + JAPANESE_TEXT.encode("utf-16-le"),
        codecs.BOM_UTF16_BE + JAPANESE_TEXT.encode("utf-16-be"),
        codecs.BOM_UTF32_LE + JAPANESE_TEXT.encode("utf-32-le"),
        codecs.BOM_UTF32_BE + JAPANESE_TEXT.encode("utf-32-be"),
        JAPANESE_TEXT.encode("utf-8"),
    ],
    ids=["utf-8-bom", "utf-16-le", "utf-16-be", "utf-32-le", "utf-32-be", "utf-8"],
)
def test正常系_decodeメソッドがBOMつきとUTF8の本文を復号すること(text_reader, data):
    assert text_reader.decode(data) == JAPANESE_TEXT


def test正常系_decodeメソッドがUTF8で復号できる本文に文字コード判定を使わないこと(text_reader, detect):
    text_reader.decode(JAPANESE_TEXT.encode("utf-8"))

    detect.assert_not_called()


@pytest.mark.parametrize(
    ("detected", "encoding"),
    [("EUC-JP", "euc_jp"), ("UTF-16LE", "utf-16-le"), ("utf-32-be", "utf-32-be")],
)
def test正常系_decodeメソッドが信用できる判定結果の文字コードで復号すること(text_reader, detect, detected, encoding):
    detect.return_value = {"encoding": detected}

    assert text_reader.decode(JAPANESE_TEXT.encode(encoding)) == JAPANESE_TEXT


def test正常系_decodeメソッドがJISの本文をUTF8より先にJISとして復号すること(text_reader):
    assert text_reader.decode(JAPANESE_TEXT.encode("iso2022_jp")) == JAPANESE_TEXT


def test正常系_decodeメソッドがJISのエスケープに似たバイト列を含むUTF8の本文を復号すること(text_reader):
    assert text_reader.decode("価格は\x1b$5です".encode()) == "価格は\x1b$5です"


def test正常系_decodeメソッドが判定できないときCP932で復号すること(text_reader):
    assert text_reader.decode(JAPANESE_TEXT.encode("cp932")) == JAPANESE_TEXT


@pytest.mark.parametrize("detected", ["cp862", "Windows-1252", "ISO-8859-1", "ascii", "no-such-codec"])
def test正常系_decodeメソッドが1バイト文字コードと判定されたシフトJISをCP932で復号すること(
    text_reader, detect, detected
):
    detect.return_value = {"encoding": detected}

    assert text_reader.decode("foo bar 検索対象".encode("cp932")) == "foo bar 検索対象"


def test正常系_decodeメソッドがCP932で復号できない内容を判定結果の文字コードで復号すること(text_reader, detect):
    detect.return_value = {"encoding": "Windows-1252"}

    assert text_reader.decode("café".encode("cp1252")) == "café"


def test正常系_decodeメソッドが空の内容を空文字として返すこと(text_reader):
    assert text_reader.decode(b"") == ""


def test正常系_decodeメソッドが文字コード判定へ先頭の一定量だけを渡すこと(text_reader, detect):
    text_reader.decode(b"\x82\xa0" * DETECTION_SAMPLE_BYTES)

    assert len(detect.call_args.args[0]) == DETECTION_SAMPLE_BYTES


def test異常系_decodeメソッドがどの文字コードでも復号できない内容でUnreadableFileErrorを送出すること(text_reader):
    with pytest.raises(UnreadableFileError, match="文字コード"):
        text_reader.decode(b"\xff\xff\xff\x80\x81\xfd")


@pytest.mark.parametrize("text", [JAPANESE_TEXT, "foo bar 検索対象", "見積書", "abc 表"])
def test正常系_readメソッドが実際の文字コード判定でシフトJISのテキストファイルを読むこと(tmp_path, text):
    path = tmp_path / "sjis.txt"
    path.write_bytes(text.encode("cp932"))

    assert TextFileReader().read(path) == text


@pytest.mark.parametrize("encoding", ["euc_jp", "iso2022_jp"])
def test正常系_readメソッドが実際の文字コード判定でシフトJIS以外の日本語テキストを読むこと(tmp_path, encoding):
    path = tmp_path / "other.txt"
    path.write_bytes(JAPANESE_TEXT.encode(encoding))

    assert TextFileReader().read(path) == JAPANESE_TEXT


def test異常系_readメソッドが存在しないテキストファイルでOSErrorを送出すること(text_reader, tmp_path):
    with pytest.raises(OSError):
        text_reader.read(tmp_path / "missing.txt")


@pytest.fixture
def docx_path(tmp_path):
    document = docx.Document()
    document.add_paragraph("最初の段落")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "表のセル"
    nested = table.cell(0, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "入れ子の表"
    document.add_paragraph("最後の段落")
    path = tmp_path / "sample.docx"
    document.save(str(path))
    return path


@pytest.mark.parametrize("expected", ["最初の段落", "表のセル", "入れ子の表", "最後の段落"])
def test正常系_readメソッドがWord文書の段落と表のテキストを返すこと(docx_path, expected):
    assert expected in DocxFileReader().read(docx_path)


def test正常系_readメソッドがWord文書のテキストを出現順に返すこと(docx_path):
    text = DocxFileReader().read(docx_path)

    assert text.index("最初の段落") < text.index("表のセル") < text.index("最後の段落")


def test異常系_readメソッドがWordの一時ファイルでUnreadableFileErrorを送出すること():
    with pytest.raises(UnreadableFileError, match="一時ファイル"):
        DocxFileReader(open_document=MagicMock()).read(Path("~$sample.docx"))


def test異常系_readメソッドがWordの一時ファイルを開こうとしないこと():
    open_document = MagicMock()
    with contextlib.suppress(UnreadableFileError):
        DocxFileReader(open_document=open_document).read(Path("~$sample.docx"))

    open_document.assert_not_called()


def _write_garbage(path: Path) -> None:
    path.write_bytes(b"PK\x03\x04garbage")


def _write_plain_zip(path: Path) -> None:
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("a.txt", "x")


def _write_broken_xml(path: Path) -> None:
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", "<bad")


@pytest.mark.parametrize("write", [_write_garbage, _write_plain_zip, _write_broken_xml])
def test異常系_readメソッドが壊れたWord文書でUnreadableFileErrorを送出すること(tmp_path, write):
    path = tmp_path / "broken.docx"
    write(path)

    with pytest.raises(UnreadableFileError, match="Word文書として開けません"):
        DocxFileReader().read(path)


def test異常系_readメソッドがWord文書の読み取り権限エラーをOSErrorのまま送出すること():
    open_document = MagicMock(side_effect=PermissionError(13, "Permission denied"))

    with pytest.raises(PermissionError):
        DocxFileReader(open_document=open_document).read(Path("locked.docx"))


def test正常系_default_readersメソッドが対応する拡張子を返すこと():
    assert set(default_readers()) == {".docx", ".txt"}
